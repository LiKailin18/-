"""文档解析：PDF（PyMuPDF）与 DOCX（python-docx）。

纯函数、无 IO 之外的副作用，方便单测（CLAUDE.md §5.1）。
解析在 BackgroundTasks 里通过线程池调用 —— 它是 CPU 密集操作，不能阻塞事件循环。

扫描版 PDF 的判定（docs/api.md §5）：平均每页可提取字符数低于阈值时，
认为没有文本层，标记为 unsupported（不抛错，交由前端提示「暂不支持 OCR」）。
"""

import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf  # PyMuPDF（1.24+ 推荐用 pymupdf 而不是已弃用的 fitz 别名）
from docx import Document

# 平均每页少于这么多字符，就认为没有文本层（扫描件）
SCANNED_CHARS_PER_PAGE = 20

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MULTI_BLANK_LINES = re.compile(r"\n{3,}")
_TRAILING_SPACES = re.compile(r"[ \t]+\n")


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    page_count: int | None
    char_count: int
    is_scanned: bool


class DocumentParseError(Exception):
    """解析失败（加密、损坏、格式异常）。"""


def clean_text(raw: str) -> str:
    """清洗：去控制字符、行尾空白、连续空行。"""
    text = _CONTROL_CHARS.sub("", raw)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _TRAILING_SPACES.sub("\n", text)
    text = _MULTI_BLANK_LINES.sub("\n\n", text)
    return text.strip()


def parse_document(path: Path, ext: str) -> ParsedDocument:
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    raise DocumentParseError(f"不支持的文件类型: {ext}")


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        with pymupdf.open(path) as doc:
            if doc.is_encrypted:
                raise DocumentParseError("PDF 已加密，无法解析")
            page_count = doc.page_count
            pages = [page.get_text("text") for page in doc]
    except DocumentParseError:
        raise
    except Exception as exc:  # PyMuPDF 的异常类型不稳定
        raise DocumentParseError(f"PDF 解析失败: {type(exc).__name__}") from exc

    text = clean_text("\n".join(pages))
    char_count = len(text)
    is_scanned = page_count > 0 and char_count / page_count < SCANNED_CHARS_PER_PAGE
    return ParsedDocument(
        text=text,
        page_count=page_count,
        char_count=char_count,
        is_scanned=is_scanned,
    )


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        document = Document(str(path))
    except Exception as exc:  # python-docx 会抛各种 zip/xml 异常
        raise DocumentParseError(f"DOCX 解析失败: {type(exc).__name__}") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # 表格内容也提取，否则资料里的表格会整段丢失
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = clean_text("\n".join(parts))
    char_count = len(text)
    return ParsedDocument(
        text=text,
        page_count=None,  # DOCX 没有固定页概念
        char_count=char_count,
        is_scanned=char_count == 0,
    )
